from io import BytesIO
import json
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from local_rag import documents
from local_rag.config import Settings
from local_rag.embeddings import DeterministicEmbedder
from local_rag.indexing import Indexer
from format_fixtures import docx_bytes, pdf_bytes


class FormatTests(unittest.TestCase):
    def setUp(self):
        self.assertTrue(hasattr(documents, "DocumentSegment"), "Structured PDF/Word extraction is not implemented")
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.notes = self.root / "notes"
        self.notes.mkdir()
        self.config = self.root / "config.json"
        self.config.write_text(json.dumps({"data_dir": "data", "sources": [{"id": "notes", "path": "notes"}], "max_chars": 32}))

    def app(self):
        return Indexer(Settings.load(self.config), embedder=DeterministicEmbedder())

    def read(self, name, raw):
        path = self.notes / name
        path.write_bytes(raw)
        return documents.read_document(path, 5 * 1024 * 1024)

    def chunks(self, document, max_chars=100):
        return documents.chunk_document(document, markdown=False, max_chars=max_chars)

    def test_pdf_page_locators_survive_repeated_passages_and_blank_pages(self):
        doc = self.read("pages.pdf", pdf_bytes(["Repeated passage", "", "Repeated passage"]))
        chunks = self.chunks(doc)
        self.assertEqual([c.locator["page"] for c in chunks], [1, 3])
        self.assertEqual([c.locator["kind"] for c in chunks], ["pdf", "pdf"])
        self.assertNotEqual(chunks[0].start_offset, chunks[1].start_offset)
        for c in chunks:
            self.assertEqual(c.text, doc.text[c.start_offset:c.end_offset])

    def test_pdf_chunks_never_cross_pages(self):
        doc = self.read("pages.pdf", pdf_bytes(["First page " * 10, "Second page " * 10]))
        chunks = self.chunks(doc, max_chars=24)
        self.assertTrue(all(len(c.text) <= 24 for c in chunks))
        self.assertEqual({c.locator["page"] for c in chunks}, {1, 2})
        self.assertTrue(all(not ("First" in c.text and "Second" in c.text) for c in chunks))

    def test_pdf_page_labels_are_separate_from_physical_page_numbers(self):
        from pypdf import PdfReader, PdfWriter
        from pypdf.constants import PageLabelStyle
        writer = PdfWriter()
        writer.clone_document_from_reader(PdfReader(BytesIO(pdf_bytes(["Preface", "Body"]))))
        writer.set_page_label(0, 0, style=PageLabelStyle.LOWERCASE_ROMAN)
        writer.set_page_label(1, 1, style=PageLabelStyle.DECIMAL)
        buffer = BytesIO()
        writer.write(buffer)
        chunks = self.chunks(self.read("labels.pdf", buffer.getvalue()))
        self.assertEqual([(c.locator["page"], c.locator["page_label"]) for c in chunks], [(1, "i"), (2, "1")])

    def test_image_only_pdf_reports_ocr_required(self):
        with self.assertRaisesRegex(ValueError, "ocr_required"):
            self.read("scan.pdf", pdf_bytes([None]))

    def test_mixed_text_and_scanned_pdf_does_not_silently_drop_a_page(self):
        with self.assertRaisesRegex(ValueError, "ocr_required"):
            self.read("mixed.pdf", pdf_bytes(["Readable text", None]))

    def test_encrypted_pdf_is_explicitly_reported(self):
        with self.assertRaisesRegex(ValueError, "pdf_encrypted"):
            self.read("locked.pdf", pdf_bytes(["Secret"], encrypted=True))

    def test_corrupt_pdf_reports_format_error(self):
        with self.assertRaisesRegex(ValueError, "invalid_pdf"):
            self.read("bad.pdf", b"%PDF-not-a-valid-document")

    def test_truly_blank_pdf_is_empty_content(self):
        self.assertEqual(self.chunks(self.read("blank.pdf", pdf_bytes([""]))), [])

    def test_docx_headings_and_blank_paragraphs_preserve_paragraph_numbers(self):
        def build(document):
            document.add_heading("Main", level=1)
            document.add_paragraph("")
            document.add_paragraph("Repeated passage")
            document.add_heading("Detail", level=2)
            document.add_paragraph("Repeated passage")
        doc = self.read("notes.docx", docx_bytes(build))
        chunks = self.chunks(doc)
        self.assertEqual([c.locator["paragraph"] for c in chunks], [1, 3, 4, 5])
        self.assertEqual(chunks[-1].heading_path, ("Main", "Detail"))
        for c in chunks:
            self.assertEqual(c.locator["kind"], "docx")
            self.assertNotIn("page", c.locator)
            self.assertEqual(c.text, doc.text[c.start_offset:c.end_offset])

    def test_docx_tables_are_in_body_order_and_cite_cells(self):
        def build(document):
            document.add_paragraph("Before")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Left"
            table.cell(0, 1).text = "Right"
            document.add_paragraph("After")
        chunks = self.chunks(self.read("table.docx", docx_bytes(build)))
        self.assertEqual([c.text for c in chunks], ["Before", "Left", "Right", "After"])
        self.assertEqual((chunks[1].locator["table"], chunks[1].locator["row"], chunks[1].locator["cell"]), (1, 1, 1))
        self.assertEqual(chunks[2].locator["cell"], 2)
        self.assertEqual(chunks[-1].locator["paragraph"], 2)

    def test_docx_nested_tables_retain_container_locations(self):
        def build(document):
            outer = document.add_table(rows=1, cols=1)
            outer.cell(0, 0).text = "Outer"
            inner = outer.cell(0, 0).add_table(rows=1, cols=1)
            inner.cell(0, 0).text = "Nested"
        chunks = self.chunks(self.read("nested.docx", docx_bytes(build)))
        nested = next(c for c in chunks if c.text == "Nested")
        self.assertEqual(len(nested.locator["table_path"]), 2)

    def test_docx_merged_cells_do_not_duplicate_text(self):
        def build(document):
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged"
        chunks = self.chunks(self.read("merged.docx", docx_bytes(build)))
        self.assertEqual([c.text for c in chunks], ["Merged"])

    def test_docx_custom_style_inherits_heading_level(self):
        def build(document):
            from docx.enum.style import WD_STYLE_TYPE
            style = document.styles.add_style("Custom section", WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = document.styles["Heading 2"]
            document.add_heading("Main", level=1)
            document.add_paragraph("Custom", style=style)
            document.add_paragraph("Body")
        chunks = self.chunks(self.read("custom.docx", docx_bytes(build)))
        self.assertEqual(chunks[-1].heading_path, ("Main", "Custom"))

    def test_docx_hyperlink_display_text_is_preserved_without_fetching_url(self):
        def build(document):
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            from docx.opc.constants import RELATIONSHIP_TYPE
            paragraph = document.add_paragraph("See ")
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("r:id"), paragraph.part.relate_to("https://example.invalid/private", RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
            run, text = OxmlElement("w:r"), OxmlElement("w:t")
            text.text = "the reference"
            run.append(text)
            hyperlink.append(run)
            paragraph._p.append(hyperlink)
        self.assertEqual(self.chunks(self.read("link.docx", docx_bytes(build)))[0].text, "See the reference")

    def test_docx_long_paragraph_splits_keep_the_same_paragraph_locator(self):
        raw = docx_bytes(lambda d: d.add_paragraph("word " * 40))
        doc = self.read("long.docx", raw)
        chunks = self.chunks(doc, max_chars=24)
        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(c.locator["paragraph"] == 1 for c in chunks))
        self.assertTrue(all(len(c.text) <= 24 for c in chunks))

    def test_docx_unicode_text_is_preserved(self):
        doc = self.read("unicode.docx", docx_bytes(lambda d: d.add_paragraph("Sørens noter — 日本語")))
        self.assertEqual(self.chunks(doc)[0].text, "Sørens noter — 日本語")

    def test_corrupt_docx_reports_format_error(self):
        with self.assertRaisesRegex(ValueError, "invalid_docx"):
            self.read("bad.docx", b"not a zip archive")

    def test_legacy_word_is_reported_as_unsupported(self):
        with self.assertRaisesRegex(ValueError, "unsupported_format"):
            self.read("old.doc", b"legacy Word file")

    def test_docx_expansion_limit_is_checked_before_parsing(self):
        from local_rag import extractors
        raw = docx_bytes(lambda d: d.add_paragraph("text"))
        with patch.object(extractors, "MAX_DOCX_UNCOMPRESSED_BYTES", 10):
            with self.assertRaisesRegex(ValueError, "document_limit"):
                self.read("large.docx", raw)

    def test_pdf_page_limit_is_checked_before_extraction(self):
        from local_rag import extractors
        with patch.object(extractors, "MAX_PDF_PAGES", 1):
            with self.assertRaisesRegex(ValueError, "document_limit"):
                self.read("many.pdf", pdf_bytes(["one", "two"]))

    def test_unsupported_docx_textbox_is_not_silently_omitted(self):
        def build(document):
            from docx.oxml import OxmlElement
            paragraph = document.add_paragraph("Normal text")
            textbox = OxmlElement("w:txbxContent")
            paragraph._p.append(textbox)
        with self.assertRaisesRegex(ValueError, "unsupported_docx_content"):
            self.read("textbox.docx", docx_bytes(build))

    def test_empty_docx_has_no_chunks(self):
        self.assertEqual(self.chunks(self.read("empty.docx", docx_bytes())), [])

    def test_pdf_and_docx_locators_persist_across_close_and_reopen(self):
        (self.notes / "file.pdf").write_bytes(pdf_bytes(["First", "Second"]))
        (self.notes / "file.docx").write_bytes(docx_bytes(lambda d: d.add_paragraph("Word body")))
        with self.app() as app:
            self.assertEqual((app.index().indexed, len(app.store.list_documents())), (2, 2))
        with self.app() as app:
            self.assertEqual([c.locator["page"] for c in app.store.chunks("notes", "file.pdf")], [1, 2])
            self.assertEqual(app.store.chunks("notes", "file.docx")[0].locator["paragraph"], 1)

    def test_pdf_extraction_failure_preserves_old_text_and_revision(self):
        path = self.notes / "file.pdf"
        path.write_bytes(pdf_bytes(["Old indexed text"]))
        with self.app() as app:
            self.assertEqual(app.index().indexed, 1)
            old = app.store.list_documents()[0].revision
            path.write_bytes(pdf_bytes([None]))
            report = app.index()
            self.assertEqual((report.failed, report.deleted), (1, 0))
            self.assertEqual(report.errors[0]["error_code"], "ocr_required")
            self.assertEqual(app.store.list_documents()[0].revision, old)
            self.assertIn("Old indexed text", app.store.chunks("notes", "file.pdf")[0].text)

    def test_corrupt_docx_update_preserves_old_revision(self):
        path = self.notes / "file.docx"
        path.write_bytes(docx_bytes(lambda d: d.add_paragraph("Old text")))
        with self.app() as app:
            self.assertEqual(app.index().indexed, 1)
            old = app.store.list_documents()[0].revision
            path.write_bytes(b"corrupt archive")
            self.assertEqual(app.index().failed, 1)
            self.assertEqual(app.store.list_documents()[0].revision, old)

    def test_empty_docx_update_clears_old_chunks(self):
        path = self.notes / "file.docx"
        path.write_bytes(docx_bytes(lambda d: d.add_paragraph("Old text")))
        with self.app() as app:
            self.assertEqual(app.index().indexed, 1)
            path.write_bytes(docx_bytes())
            self.assertEqual(app.index().indexed, 1)
            self.assertEqual(app.store.chunks("notes", "file.docx"), [])

    def test_mixed_format_indexing_runs_without_network_or_process_creation(self):
        import socket
        (self.notes / "file.pdf").write_bytes(pdf_bytes(["PDF text"]))
        (self.notes / "file.docx").write_bytes(docx_bytes(lambda d: d.add_paragraph("Word text")))
        with patch.object(socket, "socket", side_effect=AssertionError("network forbidden")), patch("subprocess.Popen", side_effect=AssertionError("process forbidden")):
            with self.app() as app:
                self.assertEqual(app.index().indexed, 2)

    def test_schema_one_upgrade_preserves_existing_indexed_content(self):
        (self.notes / "legacy.txt").write_text("Existing text")
        with self.app() as app:
            self.assertEqual(app.index().indexed, 1)
        database = self.root / "data" / "index.sqlite3"
        with sqlite3.connect(database) as connection:
            columns = [r[1] for r in connection.execute("PRAGMA table_info(chunks)")]
            if "locator" in columns:
                connection.execute("ALTER TABLE chunks DROP COLUMN locator")
            connection.execute("UPDATE meta SET value='1' WHERE key='schema_version'")
        connection.close()
        with self.app() as app:
            chunks = app.store.chunks("notes", "legacy.txt")
            self.assertEqual(chunks[0].text, "Existing text")
            self.assertEqual(chunks[0].locator["kind"], "text")
            version = app.store.connection.execute("SELECT value FROM meta WHERE key='schema_version'").fetchone()[0]
            self.assertEqual(version, "2")

    def test_docx_header_and_footer_text_have_section_locators(self):
        def build(document):
            document.add_paragraph("Body text")
            document.sections[0].header.paragraphs[0].text = "Header text"
            document.sections[0].footer.paragraphs[0].text = "Footer text"
        chunks = self.chunks(self.read("stories.docx", docx_bytes(build)))
        self.assertEqual([c.text for c in chunks], ["Body text", "Header text", "Footer text"])
        self.assertEqual(chunks[1].locator["part"], "header")
        self.assertEqual((chunks[1].locator["section"], chunks[1].locator["paragraph"]), (1, 1))
        self.assertEqual(chunks[2].locator["part"], "footer")

    def test_docx_pending_footnotes_are_explicitly_unsupported(self):
        def build(document):
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            paragraph = document.add_paragraph("Text")
            reference = OxmlElement("w:footnoteReference")
            reference.set(qn("w:id"), "1")
            paragraph.add_run()._r.append(reference)
        with self.assertRaisesRegex(ValueError, "unsupported_docx_content"):
            self.read("footnote.docx", docx_bytes(build))

    def test_docx_inline_equations_are_explicitly_unsupported(self):
        def build(document):
            from docx.oxml import OxmlElement
            document.add_paragraph("Formula")._p.append(OxmlElement("m:oMath"))
        with self.assertRaisesRegex(ValueError, "unsupported_docx_content"):
            self.read("math.docx", docx_bytes(build))

    def test_pdf_graphics_without_text_are_reported_instead_of_erasing_content(self):
        from reportlab.pdfgen import canvas
        buffer = BytesIO()
        document = canvas.Canvas(buffer, invariant=True)
        document.rect(10, 10, 100, 100)
        document.save()
        with self.assertRaisesRegex(ValueError, "pdf_no_text"):
            self.read("drawing.pdf", buffer.getvalue())

    def test_image_only_docx_reports_ocr_required(self):
        def build(document):
            from PIL import Image
            image = BytesIO()
            Image.new("RGB", (20, 20), "black").save(image, format="PNG")
            image.seek(0)
            document.add_picture(image)
        with self.assertRaisesRegex(ValueError, "ocr_required"):
            self.read("scan.docx", docx_bytes(build))

    def test_cli_inspect_exposes_pdf_and_word_locators(self):
        import os
        import subprocess
        import sys
        (self.notes / "file.pdf").write_bytes(pdf_bytes(["PDF text"]))
        (self.notes / "file.docx").write_bytes(docx_bytes(lambda d: d.add_paragraph("Word text")))
        env = dict(os.environ, PYTHONPATH=str(Path(__file__).resolve().parents[1] / "src"), PYTHONIOENCODING="utf-8")
        command = [sys.executable, "-m", "local_rag", "--config", str(self.config)]
        result = subprocess.run([*command, "index", "--test-embeddings"], env=env, capture_output=True, text=True, timeout=15)
        self.assertEqual(result.returncode, 0, result.stderr)
        for path, field in (("file.pdf", "page"), ("file.docx", "paragraph")):
            result = subprocess.run([*command, "inspect", "--source", "notes", "--path", path], env=env, capture_output=True, text=True, timeout=15)
            self.assertEqual(result.returncode, 0, result.stderr)
            data = json.loads(result.stdout)
            self.assertEqual(data["chunks"][0]["locator"][field], 1)
            self.assertFalse(data["stale"])

    def test_existing_format_profile_requires_explicit_rebuild(self):
        (self.notes / "legacy.txt").write_text("Legacy content")
        with self.app() as app:
            app.index()
            old_revision = app.store.list_documents()[0].revision
            app.store.connection.execute("UPDATE generations SET fingerprint='older-extractor' WHERE status='active'")
            from local_rag.indexing import ProfileMismatch
            with self.assertRaises(ProfileMismatch):
                app.index()
            self.assertEqual(app.store.list_documents()[0].revision, old_revision)
            self.assertEqual(app.index(rebuild=True).failed, 0)

    def test_format_pilot_verifies_a_nonempty_synthetic_ground_truth_collection(self):
        import os
        import subprocess
        import sys
        root = Path(__file__).resolve().parents[1]
        env = dict(os.environ, PYTHONPATH=str(root / "src"), PYTHONIOENCODING="utf-8")
        result = subprocess.run([sys.executable, str(root / "tools" / "format_pilot.py")], env=env,
                                capture_output=True, text=True, timeout=15)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(json.loads(result.stdout)["verified_questions"], 8)
        self.assertEqual(json.loads(result.stdout)["documents"], 4)

    def test_ole_word_container_is_reported_as_encrypted_or_legacy(self):
        with self.assertRaisesRegex(ValueError, "unsupported_docx_container"):
            self.read("protected.docx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 128)

    def test_image_only_word_body_with_header_still_requires_ocr(self):
        def build(document):
            from PIL import Image
            image = BytesIO()
            Image.new("RGB", (20, 20), "black").save(image, format="PNG")
            image.seek(0)
            document.add_picture(image)
            document.sections[0].header.paragraphs[0].text = "Scanner header"
        with self.assertRaisesRegex(ValueError, "ocr_required"):
            self.read("scan-header.docx", docx_bytes(build))

    def test_encrypted_pdf_is_rejected_before_password_or_crypto_processing(self):
        from pypdf import PdfReader
        raw = pdf_bytes(["Encrypted content"], encrypted=True)
        with patch.object(PdfReader, "_handle_encryption", side_effect=AssertionError("decryption forbidden")):
            with self.assertRaisesRegex(ValueError, "pdf_encrypted"):
                self.read("locked.pdf", raw)

    def test_cli_unicode_word_output_survives_an_ascii_console(self):
        import os
        import subprocess
        import sys
        (self.notes / "unicode.docx").write_bytes(docx_bytes(lambda d: d.add_paragraph("Søren 日本語")))
        env = dict(os.environ, PYTHONPATH=str(Path(__file__).resolve().parents[1] / "src"), PYTHONIOENCODING="ascii")
        command = [sys.executable, "-m", "local_rag", "--config", str(self.config)]
        indexed = subprocess.run([*command, "index", "--test-embeddings"], env=env, capture_output=True, timeout=15)
        self.assertEqual(indexed.returncode, 0, indexed.stderr)
        result = subprocess.run([*command, "inspect", "--source", "notes", "--path", "unicode.docx"], env=env, capture_output=True, timeout=15)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(json.loads(result.stdout)["chunks"][0]["text"], "Søren 日本語")


if __name__ == "__main__":
    unittest.main()
